"""File validation, text extraction, normalisation and chunking. Uploaded files are only ever parsed, never executed."""
import csv
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from pathlib import Path

ALLOWED = {
    "pdf": {"application/pdf"},
    "docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"},
    "xlsx": {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/octet-stream"},
    "txt": {"text/plain"},
    "csv": {"text/csv", "text/plain", "application/vnd.ms-excel"},
}
CANONICAL_MIME = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "txt": "text/plain",
    "csv": "text/csv",
}
MAX_EXTRACTED_CHARS = 5_000_000


class ValidationError(ValueError):
    """Raised when an upload fails validation. Messages are safe to show to end users."""


@dataclass
class Page:
    number: int | None
    text: str


@dataclass
class ChunkData:
    index: int
    page: int | None
    text: str


def safe_filename(name: str | None) -> str:
    """Keep only the final path component; the name is metadata only and never used to build a path."""
    base = re.split(r"[\\/]", name or "")[-1]
    base = "".join(c for c in base if c.isprintable()).strip()
    return base[:255] or "unnamed"


def get_extension(filename: str) -> str:
    return Path(filename).suffix.lower().lstrip(".")


def validate_upload(filename: str, declared_mime: str | None, data: bytes, max_bytes: int) -> str:
    """Return the validated extension or raise ValidationError. Checks size, extension, MIME and real structure."""
    if not data:
        raise ValidationError("The file is empty.")
    if len(data) > max_bytes:
        raise ValidationError(f"The file exceeds the {max_bytes // (1024 * 1024)} MB limit.")
    ext = get_extension(filename)
    if ext not in ALLOWED:
        raise ValidationError(f"Unsupported file type '.{ext}'. Allowed: {', '.join(sorted(ALLOWED))}.")
    mime = (declared_mime or "").split(";")[0].strip().lower()
    if mime and mime not in ALLOWED[ext]:
        raise ValidationError(f"Declared content type '{mime}' does not match a .{ext} file.")
    if ext == "pdf":
        if not data.lstrip()[:5] == b"%PDF-":
            raise ValidationError("The file is not a valid PDF (missing PDF header).")
    elif ext in ("docx", "xlsx"):
        _validate_ooxml(data, ext)
    else:
        if b"\x00" in data[:8192]:
            raise ValidationError("The file contains binary data and is not valid text.")
        try:
            data[:65536].decode("utf-8")
        except UnicodeDecodeError as exc:
            if exc.start < len(data[:65536]) - 4:  # ignore a multibyte char split at the sample boundary
                raise ValidationError("The file is not valid UTF-8 text.") from exc
    return ext


def _validate_ooxml(data: bytes, ext: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = set(zf.namelist())
            total = sum(i.file_size for i in zf.infolist())
            if total > 200 * 1024 * 1024:
                raise ValidationError("The archive expands to an unreasonable size.")
    except zipfile.BadZipFile as exc:
        raise ValidationError(f"The file is not a valid .{ext} container.") from exc
    required = "word/document.xml" if ext == "docx" else "xl/workbook.xml"
    if "[Content_Types].xml" not in names or required not in names:
        raise ValidationError(f"The file does not have the structure of a .{ext} document.")


def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = "".join(c for c in text if c in "\n\t" or unicodedata.category(c)[0] != "C")
    text = re.sub(r"[ \t\f\v]+", " ", text)
    text = re.sub(r" ?\n ?", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pages(path: Path, ext: str) -> list[Page]:
    if ext == "pdf":
        pages = _extract_pdf(path)
    elif ext == "docx":
        pages = _extract_docx(path)
    elif ext == "xlsx":
        pages = _extract_xlsx(path)
    elif ext == "csv":
        pages = _extract_csv(path)
    elif ext == "txt":
        pages = [Page(None, path.read_text(encoding="utf-8", errors="replace"))]
    else:
        raise ValidationError(f"Unsupported file type '.{ext}'.")
    cleaned = [Page(p.number, normalize_text(p.text)) for p in pages]
    cleaned = [p for p in cleaned if p.text]
    if sum(len(p.text) for p in cleaned) > MAX_EXTRACTED_CHARS:
        raise ValidationError("The extracted text is too large to process.")
    return cleaned


def _extract_pdf(path: Path) -> list[Page]:
    import pymupdf

    pages: list[Page] = []
    with pymupdf.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            pages.append(Page(i, page.get_text("text")))
    return pages


def _extract_docx(path: Path) -> list[Page]:
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return [Page(None, "\n".join(parts))]  # DOCX has no stable page numbers


def _extract_xlsx(path: Path) -> list[Page]:
    import openpyxl

    pages = []
    with path.open("rb") as fh:  # stream: stored files have no extension, which openpyxl would reject
        wb = openpyxl.load_workbook(fh, read_only=True, data_only=True)
        for ws in wb.worksheets:
            lines = [f"Sheet: {ws.title}"]
            for row in ws.iter_rows(values_only=True):
                if any(c is not None for c in row):
                    lines.append(" | ".join("" if c is None else str(c) for c in row))
            pages.append(Page(None, "\n".join(lines)))
        wb.close()
    return pages


def _extract_csv(path: Path) -> list[Page]:
    text = path.read_text(encoding="utf-8", errors="replace")
    rows = csv.reader(io.StringIO(text))
    return [Page(None, "\n".join(" | ".join(r) for r in rows))]


def chunk_pages(pages: list[Page], size: int = 900, overlap: int = 120) -> list[ChunkData]:
    """Split text into overlapping chunks, preferring paragraph/sentence boundaries. Chunks never span pages."""
    if size <= overlap:
        raise ValueError("chunk size must exceed overlap")
    chunks: list[ChunkData] = []
    for page in pages:
        text, start = page.text, 0
        while start < len(text):
            end = min(start + size, len(text))
            if end < len(text):
                window = text[start:end]
                cut = max(window.rfind("\n\n"), window.rfind(". "), window.rfind("\n"))
                if cut > size * 0.5:
                    end = start + cut + 1
            piece = text[start:end].strip()
            if piece:
                chunks.append(ChunkData(len(chunks), page.number, piece))
            if end >= len(text):
                break
            start = max(end - overlap, start + 1)
    return chunks
