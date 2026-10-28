import os
import tempfile
from pathlib import Path

# Configure BEFORE importing the app: in-memory DB and vector store, hashing embedder, no Ollama required.
_TMP = tempfile.mkdtemp(prefix="infomind-test-")
os.environ.update({
    "ENVIRONMENT": "test", "DATABASE_URL": "sqlite://", "QDRANT_URL": ":memory:", "EMBEDDING_PROVIDER": "hash",
    "EMBEDDING_DIMENSIONS": "256", "RUN_WORKER_IN_API": "false", "STORAGE_DIR": str(Path(_TMP) / "uploads"),
    "MODEL_DIR": str(Path(_TMP) / "models"), "JWT_SECRET": "test-secret-test-secret-test-secret-123",
    "OLLAMA_BASE_URL": "http://127.0.0.1:9", "RATE_LIMIT_PER_MINUTE": "1000",
})

import io  # noqa: E402

import pytest  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402

from app.core.errors import ServiceUnavailable  # noqa: E402
from app.core.ratelimit import reset_rate_limits  # noqa: E402
from app.core.security import hash_password  # noqa: E402
from app.db.session import Base, SessionLocal, engine  # noqa: E402
from app.main import app  # noqa: E402
from app.models import User  # noqa: E402
from app.services import classification_service, document_service, rag_service  # noqa: E402
from app.services.vector_store import _client  # noqa: E402

PASSWORD = "correct-horse-battery"


class FakeLLM:
    """Scriptable stand-in for Ollama. `responder(prompt) -> str` returns the raw model output."""
    name = "fake-llm"

    def __init__(self):
        self.calls: list[str] = []
        self.fail = False
        self.responder = self.default

    @staticmethod
    def default(prompt: str) -> str:
        import json
        return json.dumps({"supported": True, "citations": ["E1"], "answer": "Per the evidence, the stated fact holds.",
                           "inference": ""})

    def chat_json(self, system: str, user: str) -> str:
        self.calls.append(user)
        if self.fail:
            raise ServiceUnavailable("Ollama is not reachable.")
        return self.responder(user)

    def chat_json_stream(self, system: str, user: str):
        raw = self.chat_json(system, user)
        for i in range(0, len(raw), 7):  # small odd-sized pieces, so escapes and keys get split across chunks
            yield raw[i:i + 7]


@pytest.fixture(autouse=True)
def clean_state():
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    client = _client()
    for c in client.get_collections().collections:
        client.delete_collection(c.name)
    classification_service._cache.clear()
    reset_rate_limits()
    yield


@pytest.fixture
def llm(monkeypatch) -> FakeLLM:
    fake = FakeLLM()
    monkeypatch.setattr(rag_service, "get_llm", lambda: fake)
    return fake


@pytest.fixture
def client() -> TestClient:
    return TestClient(app, raise_server_exceptions=False)


@pytest.fixture
def db():
    s = SessionLocal()
    yield s
    s.close()


def make_user(db, email: str, role: str = "analyst") -> User:
    u = User(email=email, full_name=email.split("@")[0], password_hash=hash_password(PASSWORD), role=role)
    db.add(u)
    db.commit()
    return u


def login(client: TestClient, email: str) -> dict:
    r = client.post("/api/v1/auth/login", json={"email": email, "password": PASSWORD})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture
def make_auth(client, db):
    def _make(email: str, role: str = "analyst") -> dict:
        make_user(db, email, role)
        return login(client, email)
    return _make


@pytest.fixture
def analyst(make_auth):
    return make_auth("analyst@example.com", "analyst")


@pytest.fixture
def other_analyst(make_auth):
    return make_auth("other@example.com", "analyst")


@pytest.fixture
def admin(make_auth):
    return make_auth("admin@example.com", "admin")


@pytest.fixture
def reviewer(make_auth):
    return make_auth("reviewer@example.com", "reviewer")


def process_jobs() -> int:
    with SessionLocal() as s:
        return document_service.run_pending_jobs(s)


@pytest.fixture
def upload(client):
    def _upload(headers, name="notes.txt", content: bytes | str = b"hello", mime=None):
        data = content.encode() if isinstance(content, str) else content
        import mimetypes
        mime = mime or mimetypes.guess_type(name)[0] or "application/octet-stream"
        r = client.post("/api/v1/documents", headers=headers, files=[("files", (name, data, mime))])
        return r
    return _upload


@pytest.fixture
def upload_and_process(upload):
    def _go(headers, name, content, mime=None):
        r = upload(headers, name, content, mime)
        assert r.status_code == 201, r.text
        process_jobs()
        return r.json()["results"][0]["document"]["id"]
    return _go


# ---- synthetic file builders
def make_pdf(text_pages: list[str]) -> bytes:
    import pymupdf
    doc = pymupdf.open()
    for t in text_pages:
        page = doc.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 780), t, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def make_docx(paragraphs: list[str]) -> bytes:
    import docx
    d = docx.Document()
    for p in paragraphs:
        d.add_paragraph(p)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx(rows: list[list]) -> bytes:
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


SAMPLE_TEXT = ("The purchasing department issues the purchase order and negotiates supplier quotations. "
               "The annual budget for fleet maintenance is 250000 dollars, approved on 2025-03-14. "
               "Invoice payment terms are thirty days after delivery.")
